from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from textwrap import wrap

from django.core.files.base import ContentFile
from django.core.management import call_command
from django.core.exceptions import ValidationError
from django.core.exceptions import ImproperlyConfigured
from django.forms import models
from django.shortcuts import get_object_or_404, render
from django.db.models import Count, Q
from django.views import View
from rest_framework.response import Response
from rest_framework import generics, viewsets, status

from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import IsAuthenticated

# Create your views here.

import csv
import hashlib
import hmac
import logging
import re
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone as dt_timezone

from django.contrib.auth import get_user_model
from django.db.models import Sum
from django.db.models.functions import TruncMonth
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
import json

from django.conf import settings
from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.csrf import csrf_exempt
import requests
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


from campaigns.serializers import (
    BatchLeadSerializer,
    BatchSerializer,
    CitySerializer,
    CompanyContactSerializer,
    CompanySerializer,
    CountrySerializer,
    MessageLogSerializer,
    SequenceSerializer,
    SequenceStepSerializer,
    StateSerializer,
    MessageTemplateSerializer,
    LeadSerializer,
    OpportunityCreateSerializer,
    OpportunityModuleSerializer,
    QuotationSerializer,
    SalesOrderSerializer,
    FinalOfferUpdateSerializer,
    TaskSerializer,
)

from .models import (
    Batch,
    BatchLead,
    Channel,
    City,
    Company,
    CompanyContact,
    Country,
    EmailSuppression,
    MessageLog,
    Opportunity,
    OpportunityModule,
    Quotation,
    QuotationItem,
    SalesOrder,
    Sequence,
    SequenceStep,
    State,
    MessageTemplate,
    Lead,
    Task,
    TaskAttachment,
)
from campaigns.unsubscribe import (
    is_email_suppressed,
    record_unsubscribe,
    resolve_unsubscribe_token,
    sync_mailgun_suppression,
)


logger = logging.getLogger(__name__)


def _parse_import_date(value):
    value = (value or "").strip()
    if not value:
        return None

    parsed = parse_date(value)
    if parsed is not None:
        return parsed.isoformat()

    for date_format in ("%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, date_format).date().isoformat()
        except ValueError:
            continue
    return value


def _spreadsheet_cell_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "")


# ------------------------------------- masters ----------------------


# company


class CountryViewSet(viewsets.ModelViewSet):
    queryset = Country.objects.all()
    serializer_class = CountrySerializer


class StateViewSet(viewsets.ModelViewSet):
    queryset = State.objects.all()
    serializer_class = StateSerializer


class CityViewSet(viewsets.ModelViewSet):
    queryset = City.objects.all()
    serializer_class = CitySerializer


class CompanyListView(APIView):
    def get(self, request):
        companies = (
            Company.objects.filter(is_active=True)
            .annotate(
                contacts_count=Count("contacts", filter=Q(contacts__is_active=True))
            )
            .order_by("name")
        )
        return Response(
            {
                "count": companies.count(),
                "results": CompanySerializer(companies, many=True).data,
            }
        )

    def post(self, request):
        serializer = CompanySerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        company = serializer.save()
        company.contacts_count = 0
        return Response(CompanySerializer(company).data, status=status.HTTP_201_CREATED)


class CompanyDetailView(generics.RetrieveUpdateAPIView):
    queryset = Company.objects.filter(is_active=True).annotate(
        contacts_count=Count("contacts", filter=Q(contacts__is_active=True))
    )
    serializer_class = CompanySerializer


class CompanyContactListCreateView(APIView):
    def get(self, request, company_id):
        company = get_object_or_404(Company, pk=company_id, is_active=True)
        contacts = company.contacts.filter(is_active=True).order_by(
            "-is_primary", "name"
        )
        return Response(CompanyContactSerializer(contacts, many=True).data)

    def post(self, request, company_id):
        company = get_object_or_404(Company, pk=company_id, is_active=True)
        serializer = CompanyContactSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        return Response(
            CompanyContactSerializer(serializer.save(company=company)).data,
            status=status.HTTP_201_CREATED,
        )


class CompanyContactDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = CompanyContactSerializer

    def get_queryset(self):
        return CompanyContact.objects.filter(
            company_id=self.kwargs["company_id"],
            company__is_active=True,
            is_active=True,
        )

    def perform_destroy(self, instance):
        instance.is_active = False
        instance.save(update_fields=["is_active"])


def _task_queryset_for_user(user):
    queryset = (
        Task.objects.filter(is_deleted=False)
        .select_related("assigned_to", "created_by", "company", "parent_task")
        .prefetch_related("attachments", "recurring_instances")
    )
    if not (user.is_superuser or user.is_staff):
        queryset = queryset.filter(Q(assigned_to=user) | Q(created_by=user))
    return queryset


def _add_task_attachments(task, request):
    files = request.FILES.getlist("files")
    names = (
        request.data.getlist("file_names") if hasattr(request.data, "getlist") else []
    )
    for index, file in enumerate(files):
        TaskAttachment.objects.create(
            task=task,
            file=file,
            file_name=(
                names[index] if index < len(names) and names[index] else file.name
            ),
        )


class TaskListCreateView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get(self, request):
        tasks = _task_queryset_for_user(request.user)
        for field in ("module", "status", "reference_id"):
            value = request.query_params.get(field)
            if value:
                tasks = tasks.filter(**{field: value})
        assigned_to = request.query_params.get("assigned_to")
        if assigned_to:
            tasks = tasks.filter(assigned_to_id=assigned_to)
        return Response(
            TaskSerializer(tasks, many=True, context={"request": request}).data
        )

    def post(self, request):
        serializer = TaskSerializer(data=request.data, context={"request": request})
        serializer.is_valid(raise_exception=True)
        task = serializer.save(created_by=request.user)
        _add_task_attachments(task, request)
        return Response(
            TaskSerializer(task, context={"request": request}).data,
            status=status.HTTP_201_CREATED,
        )


class TaskDetailView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get_object(self, request, pk):
        return get_object_or_404(_task_queryset_for_user(request.user), pk=pk)

    def get(self, request, pk):
        task = self.get_object(request, pk)
        return Response(TaskSerializer(task, context={"request": request}).data)

    def put(self, request, pk):
        return self._update(request, pk, partial=False)

    def patch(self, request, pk):
        return self._update(request, pk, partial=True)

    def _update(self, request, pk, partial):
        task = self.get_object(request, pk)
        serializer = TaskSerializer(
            task,
            data=request.data,
            partial=partial,
            context={"request": request},
        )
        serializer.is_valid(raise_exception=True)
        previous_followup_date = task.followup_date
        previous_remarks = task.remarks
        followup_date_changed = (
            "followup_date" in serializer.validated_data
            and serializer.validated_data["followup_date"] != previous_followup_date
        )
        history = list(task.followup_history or [])
        task = serializer.save()
        if followup_date_changed:
            history.append(
                {
                    "id": len(history) + 1,
                    "followup_date": str(previous_followup_date or ""),
                    "remark": previous_remarks,
                    "updated_by": request.user.get_full_name()
                    or request.user.get_username(),
                    "timestamp": timezone.now().strftime("%d-%m-%Y %H:%M:%S"),
                }
            )
            task.followup_history = history
            task.save(update_fields=["followup_history", "updated_at"])
        _add_task_attachments(task, request)
        return Response(TaskSerializer(task, context={"request": request}).data)

    def delete(self, request, pk):
        task = self.get_object(request, pk)
        task.is_deleted = True
        task.save(update_fields=["is_deleted", "updated_at"])
        return Response(status=status.HTTP_204_NO_CONTENT)


class TaskOptionsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user_model = get_user_model()
        users = user_model.objects.filter(is_active=True).order_by(
            "first_name", "last_name"
        )
        companies = Company.objects.filter(is_active=True).order_by("name")
        return Response(
            {
                "modules": [
                    {"value": value, "label": label}
                    for value, label in Task.MODULE_CHOICES
                ],
                "priorities": [
                    {"value": value, "label": label}
                    for value, label in Task.PRIORITY_CHOICES
                ],
                "statuses": [
                    {"value": value, "label": label}
                    for value, label in Task.STATUS_CHOICES
                ],
                "repeat_frequencies": [
                    {"value": value, "label": label}
                    for value, label in Task.REPEAT_FREQUENCY_CHOICES
                ],
                "users": [
                    {
                        "value": user.pk,
                        "label": user.get_full_name() or user.get_username(),
                    }
                    for user in users
                ],
                "companies": [
                    {"value": company.pk, "label": company.name}
                    for company in companies
                ],
            }
        )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def get_task_dashboard(request):
    today = timezone.localdate()
    base_queryset = Task.objects.filter(is_deleted=False)
    user = request.user
    role_code = getattr(getattr(user, "role", None), "code", None)
    can_view_all = (
        user.is_superuser or user.is_staff or role_code in {"superadmin", "manager"}
    )
    if not can_view_all:
        base_queryset = base_queryset.filter(Q(assigned_to=user) | Q(created_by=user))

    selected_user_id = request.query_params.get("user_id")
    if selected_user_id:
        try:
            selected_user_id = Task._meta.get_field("assigned_to").target_field.to_python(
                selected_user_id
            )
        except (TypeError, ValueError, ValidationError):
            return Response(
                {"error": "Invalid user_id"}, status=status.HTTP_400_BAD_REQUEST
            )
        base_queryset = base_queryset.filter(assigned_to_id=selected_user_id)

    return Response(
        {
            "total": base_queryset.count(),
            "by_status": list(
                base_queryset.values("status")
                .annotate(count=Count("id"))
                .order_by("status")
            ),
            "by_priority": list(
                base_queryset.values("priority")
                .annotate(count=Count("id"))
                .order_by("priority")
            ),
            "by_module": list(
                base_queryset.values("module")
                .annotate(count=Count("id"))
                .order_by("module")
            ),
            "in_progress": base_queryset.filter(status="IN_PROGRESS").count(),
            "pending": base_queryset.filter(status="PENDING").count(),
            "recurring": base_queryset.filter(is_recurring=True).count(),
            "overdue": base_queryset.filter(
                followup_date__lt=today, status__in=["PENDING", "IN_PROGRESS"]
            ).count(),
            "due_today": base_queryset.filter(followup_date=today).count(),
            "upcoming": base_queryset.filter(
                followup_date__gt=today, status__in=["PENDING", "IN_PROGRESS"]
            ).count(),
            "miscellaneous": base_queryset.filter(
                task_regarding__iexact="MISCELLANEOUS"
            )
            .exclude(status="COMPLETED")
            .count(),
            "completed": base_queryset.filter(status="COMPLETED").count(),
        }
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def get_activities(request):
    today = timezone.localdate()
    next_week = today + timedelta(days=7)
    user = request.user
    role_code = getattr(getattr(user, "role", None), "code", None)
    can_view_all = (
        user.is_superuser or user.is_staff or role_code in {"superadmin", "manager"}
    )

    base_queryset = Task.objects.filter(is_deleted=False)
    if not can_view_all:
        base_queryset = base_queryset.filter(Q(assigned_to=user) | Q(created_by=user))

    task_groups = {
        "today": base_queryset.filter(followup_date=today, status="PENDING"),
        "upcoming": base_queryset.filter(
            followup_date__gt=today,
            followup_date__lte=next_week,
            status="PENDING",
        ),
        "overdue": base_queryset.filter(
            followup_date__lt=today,
            status__in=["PENDING", "IN_PROGRESS"],
        ),
        "recurring": base_queryset.filter(is_recurring=True),
    }

    response_data = {}
    for group_name, queryset in task_groups.items():
        tasks = queryset.select_related(
            "assigned_to", "created_by", "company", "parent_task"
        ).prefetch_related("attachments", "recurring_instances")
        response_data[group_name] = {
            "count": queryset.count(),
            "tasks": TaskSerializer(
                tasks, many=True, context={"request": request}
            ).data,
        }

    return Response(response_data)


class LeadCSVImportView(APIView):
    def post(self, request):
        csv_file = request.FILES.get("file")

        if not csv_file:
            return Response(
                {"error": "A CSV or Excel file is required."}, status=status.HTTP_400_BAD_REQUEST
            )

        extension = csv_file.name.lower().rsplit(".", 1)[-1]
        if extension not in {"csv", "xls", "xlsx"}:
            return Response(
                {"error": "Only CSV, XLS, or XLSX files are allowed."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            if extension == "csv":
                decoded_file = csv_file.read().decode("utf-8-sig")
                sample = decoded_file[:8192]
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.DictReader(
                    decoded_file.splitlines(),
                    dialect=dialect,
                )
                rows = (
                    {
                        (key or "").strip().lower(): (value or "")
                        for key, value in row.items()
                    }
                    for row in reader
                )
            elif extension == "xlsx":
                from openpyxl import load_workbook

                workbook = load_workbook(csv_file, read_only=True, data_only=True)
                sheet = workbook.active
                values = sheet.iter_rows(values_only=True)
                headers = [str(value or "").strip().lower() for value in next(values, ())]
                rows = (
                    {
                        header: _spreadsheet_cell_value(value)
                        for header, value in zip(headers, row)
                        if header
                    }
                    for row in values
                    if any(value is not None for value in row)
                )
            else:
                import xlrd

                workbook = xlrd.open_workbook(file_contents=csv_file.read())
                sheet = workbook.sheet_by_index(0)
                headers = [str(value or "").strip().lower() for value in sheet.row_values(0)]

                def xls_rows():
                    for row_index in range(1, sheet.nrows):
                        row = {}
                        for column_index, header in enumerate(headers):
                            if not header:
                                continue
                            cell = sheet.cell(row_index, column_index)
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(
                                    cell.value, workbook.datemode
                                ).date().isoformat()
                            else:
                                value = _spreadsheet_cell_value(cell.value)
                            row[header] = value
                        yield row

                rows = xls_rows()

            total_rows = 0
            imported = 0
            duplicates = 0
            invalid = 0
            row_errors = []
            valid_sources = {value for value, _label in Lead.SOURCE_CHOICES}
            valid_statuses = {value for value, _label in Lead.STATUS_CHOICES}

            for row in rows:
                total_rows += 1
                source = (row.get("source") or "").strip().upper()
                source = source if source in valid_sources else "OTHER"
                lead_status = (row.get("status") or "").strip().upper()
                lead_status = lead_status if lead_status in valid_statuses else "OTHER"
                data = {
                    "company_name": (row.get("company_name") or "").strip(),
                    "company_address": (row.get("company_address") or "").strip(),
                    "company_city": (row.get("company_city") or "").strip(),
                    "company_state": (row.get("company_state") or "").strip(),
                    "company_pincode": (row.get("company_pincode") or "").strip(),
                    "company_country": (row.get("company_country") or "India").strip(),
                    "company_gst_number": (row.get("company_gst_number") or "").strip(),
                    "contact_name": (row.get("contact_name") or "").strip(),
                    "contact_designation": (
                        row.get("contact_designation") or ""
                    ).strip(),
                    "contact_mobile": (row.get("contact_mobile") or "").strip(),
                    "contact_email": (row.get("contact_email") or "").strip(),
                    "enquiry_for": (row.get("enquiry_for") or "LEAD").strip().upper(),
                    "enquiry_date": _parse_import_date(row.get("enquiry_date"))
                    or timezone.localdate().isoformat(),
                    "source": source,
                    "status": lead_status,
                    "followup_date": _parse_import_date(row.get("followup_date")),
                    "followup_remarks": (row.get("followup_remarks") or "").strip(),
                    "is_opportunity": (row.get("is_opportunity") or "").strip().lower()
                    in {"1", "true", "yes", "y"},
                }
                serializer = LeadSerializer(data=data)
                if not serializer.is_valid():
                    invalid += 1
                    row_errors.append(
                        {"row": total_rows + 1, "errors": serializer.errors}
                    )
                    continue
                serializer.save()
                imported += 1

            return Response(
                {
                    "message": "Leads imported successfully.",
                    "total_rows": total_rows,
                    "imported": imported,
                    "duplicates": duplicates,
                    "invalid": invalid,
                    "row_errors": row_errors,
                },
                status=status.HTTP_201_CREATED,
            )

        except UnicodeDecodeError:
            return Response(
                {"error": "CSV must be UTF-8 encoded."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except (ValueError, IndexError, KeyError) as exc:
            return Response(
                {"error": f"The uploaded spreadsheet could not be read: {exc}"},
                status=status.HTTP_400_BAD_REQUEST,
            )


class LeadOptionsView(APIView):
    def get(self, request):
        def to_options(choices):
            return [{"value": value, "label": label} for value, label in choices]

        return Response(
            {
                "enquiry_for": to_options(Lead.ENQUIRY_FOR_CHOICES),
                "source": to_options(Lead.SOURCE_CHOICES),
                "status": to_options(Lead.STATUS_CHOICES),
            },
            status=status.HTTP_200_OK,
        )


class LeadListView(APIView):

    def get(self, request):
        leads = (
            Lead.objects.filter(is_active=True, is_opportunity=False)
            .select_related("company", "contact")
            .order_by("-enquiry_date", "-created_at")
        )
        suppressed_emails = set(
            EmailSuppression.objects.values_list("email", flat=True)
        )
        serializer = LeadSerializer(
            leads,
            many=True,
            context={"suppressed_emails": suppressed_emails},
        )

        return Response(
            {"count": leads.count(), "results": serializer.data},
            status=status.HTTP_200_OK,
        )

    def post(self, request):
        serializer = LeadSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class LeadDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = Lead.objects.filter(is_active=True).select_related("company", "contact")
    serializer_class = LeadSerializer

    def perform_destroy(self, instance):
        instance.is_active = False
        instance.save(update_fields=["is_active", "updated_at"])


class OpportunityListCreateView(generics.ListCreateAPIView):
    queryset = Opportunity.objects.filter(is_active=True).prefetch_related("modules")
    serializer_class = OpportunityCreateSerializer


class OpportunityDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = Opportunity.objects.filter(is_active=True).prefetch_related("modules")
    serializer_class = OpportunityCreateSerializer

    def perform_destroy(self, instance):
        instance.is_active = False
        instance.save(update_fields=["is_active", "updated_at"])


class OpportunityModuleListCreateView(generics.ListCreateAPIView):
    serializer_class = OpportunityModuleSerializer

    def get_opportunity(self):
        return get_object_or_404(
            Opportunity,
            pk=self.kwargs["opportunity_id"],
            is_active=True,
        )

    def get_queryset(self):
        return OpportunityModule.objects.filter(
            opportunity_id=self.kwargs["opportunity_id"],
            opportunity__is_active=True,
        ).order_by("module_number")

    def perform_create(self, serializer):
        opportunity = self.get_opportunity()
        module_number = serializer.validated_data.get("module_number")
        if module_number is None:
            last_number = (
                opportunity.modules.order_by("-module_number")
                .values_list("module_number", flat=True)
                .first()
            )
            module_number = (last_number or 0) + 1
        serializer.save(opportunity=opportunity, module_number=module_number)


class OpportunityModuleDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = OpportunityModuleSerializer

    def get_queryset(self):
        return OpportunityModule.objects.filter(
            opportunity_id=self.kwargs["opportunity_id"],
            opportunity__is_active=True,
        )


class QuotationListCreateView(generics.ListCreateAPIView):
    queryset = (
        Quotation.objects.select_related(
            "company",
            "opportunity",
            "opportunity__company",
            "opportunity__lead__company",
            "created_by",
        )
        .prefetch_related(
            "items__module",
            "company__contacts",
            "opportunity__company__contacts",
            "opportunity__lead__company__contacts",
        )
        .all()
    )
    serializer_class = QuotationSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = super().get_queryset()
        status_filter = self.request.query_params.get("status")
        opportunity_id = self.request.query_params.get("opportunity")
        search = self.request.query_params.get("search")

        if status_filter:
            queryset = queryset.filter(status=status_filter)
        if opportunity_id:
            queryset = queryset.filter(opportunity_id=opportunity_id)
        cpo_created = self.request.query_params.get("cpo_created")
        if cpo_created is not None:
            queryset = queryset.filter(cpo_created=cpo_created.lower() == "true")
        if search:
            queryset = queryset.filter(
                Q(quotation_number__icontains=search)
                | Q(opportunity__enquiry_for__icontains=search)
            )
        return queryset.order_by("-created_at")


class QuotationDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = (
        Quotation.objects.select_related(
            "company",
            "opportunity",
            "opportunity__company",
            "opportunity__lead__company",
            "created_by",
        )
        .prefetch_related(
            "items__module",
            "company__contacts",
            "opportunity__company__contacts",
            "opportunity__lead__company__contacts",
        )
        .all()
    )
    serializer_class = QuotationSerializer
    permission_classes = [IsAuthenticated]

def _generate_final_offer_pdf_bytes(quotation):
    company = quotation.company or (
        quotation.opportunity.company if quotation.opportunity else None
    )
    if company is None and quotation.opportunity:
        try:
            company = quotation.opportunity.lead.company
        except (Lead.DoesNotExist, AttributeError):
            company = None
    company_name = company.name if company else "N/A"
    city_name = company.city.name if company and company.city else "N/A"
    file_name = f"FINAL-OFFER-{quotation.quotation_number}.pdf"

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    pdf.setTitle(file_name)
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(72, height - 48, "FINAL OFFER")

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(72, height - 78, "Company:")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(160, height - 78, company_name)

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(72, height - 98, "Quotation No.:")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(160, height - 98, quotation.quotation_number)

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(72, height - 118, "FO Date:")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(
        160,
        height - 118,
        (quotation.fo_date or timezone.now().date()).strftime("%d-%m-%Y"),
    )

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(72, height - 138, "City:")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(160, height - 138, city_name)

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(72, height - 158, "Total:")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(160, height - 158, f"₹ {Decimal(quotation.total_amount or 0):,.2f}")

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(72, height - 205, "Item")
    pdf.drawString(220, height - 205, "Description")
    pdf.drawString(470, height - 205, "Amount")
    pdf.line(72, height - 212, 540, height - 212)

    y = height - 235
    for index, item in enumerate(quotation.items.all(), start=1):
        label = item.module_name or f"Item {index}"
        description = item.description or "—"
        wrapped = wrap(description, width=42)
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(72, y, label[:30])
        pdf.setFont("Helvetica", 9)
        for offset, line in enumerate(wrapped[:3]):
            pdf.drawString(220, y - offset * 12, line[:80])
        pdf.drawRightString(540, y, f"₹ {Decimal(item.amount or 0):,.2f}")
        y -= 38
        if y < 90:
            pdf.showPage()
            y = height - 60

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(72, 90, "Terms & Conditions:")
    pdf.setFont("Helvetica", 9)
    terms = quotation.terms_conditions or "Standard quotation terms apply."
    for offset, line in enumerate(wrap(terms, width=110)[:6]):
        pdf.drawString(72, 72 + (offset * -12), line)

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(72, 30, "Thank you")
    pdf.save()
    return buffer.getvalue(), file_name


class GenerateFinalOfferView(generics.GenericAPIView):
    queryset = Quotation.objects.all()
    serializer_class = QuotationSerializer
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        quotation = self.get_object()
        pdf_bytes, file_name = _generate_final_offer_pdf_bytes(quotation)
        quotation.final_offer_generated = True
        quotation.fo_date = quotation.fo_date or timezone.now().date()
        quotation.final_offer_pdf.save(file_name, ContentFile(pdf_bytes), save=False)
        quotation.save(
            update_fields=[
                "final_offer_generated",
                "fo_date",
                "final_offer_pdf",
                "updated_at",
            ]
        )

        if (
            request.headers.get("Accept", "").lower().find("application/pdf") >= 0
            or request.GET.get("download") == "1"
        ):
            response = HttpResponse(pdf_bytes, content_type="application/pdf")
            response["Content-Disposition"] = f'inline; filename="{file_name}"'
            return response

        serializer = self.get_serializer(quotation)
        return Response(serializer.data, status=status.HTTP_200_OK)


class UpdateFinalOfferView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def patch(self, request, pk):
        quotation = get_object_or_404(
            Quotation.objects.select_for_update().select_related(
                "company", "company__city", "company__state", "opportunity"
            ),
            pk=pk,
        )
        serializer = FinalOfferUpdateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        items_data = serializer.validated_data["items"]

        for item_data in items_data:
            module = item_data.get("module")
            if (
                module
                and quotation.opportunity_id
                and module.opportunity_id != quotation.opportunity_id
            ):
                return Response(
                    {
                        "items": "An item module must belong to this quotation's opportunity."
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

        quotation.items.all().delete()
        for item_data in items_data:
            item_data = dict(item_data)
            module = item_data.pop("module", None)
            module_name = item_data.pop("module_name", "")
            QuotationItem.objects.create(
                quotation=quotation,
                module=module,
                module_name=module_name or (module.module_name if module else ""),
                **item_data,
            )

        subtotal = sum(
            (item.amount for item in quotation.items.all()),
            Decimal("0.00"),
        )
        taxable_amount = max(
            subtotal - Decimal(quotation.discount or 0), Decimal("0.00")
        )
        tax_amount = (
            taxable_amount * Decimal(quotation.tax_percentage or 0) / Decimal("100")
        )
        cgst_amount = (
            tax_amount / 2 if quotation.tax_type == "INTRA_STATE" else Decimal("0.00")
        )
        sgst_amount = cgst_amount
        igst_amount = (
            tax_amount
            if quotation.tax_type and quotation.tax_type != "INTRA_STATE"
            else Decimal("0.00")
        )

        quotation.fo_date = serializer.validated_data["fo_date"]
        quotation.subtotal = subtotal
        quotation.taxable_amount = taxable_amount
        quotation.tax_amount = tax_amount
        quotation.cgst_amount = cgst_amount
        quotation.sgst_amount = sgst_amount
        quotation.igst_amount = igst_amount
        quotation.total_amount = (
            taxable_amount + tax_amount + Decimal(quotation.packing_charges or 0)
        )
        quotation.final_offer_generated = True
        pdf_bytes, file_name = _generate_final_offer_pdf_bytes(quotation)
        old_pdf_name = (
            quotation.final_offer_pdf.name if quotation.final_offer_pdf else None
        )
        quotation.final_offer_pdf.save(file_name, ContentFile(pdf_bytes), save=False)
        quotation.save(
            update_fields=[
                "fo_date",
                "subtotal",
                "taxable_amount",
                "tax_amount",
                "cgst_amount",
                "sgst_amount",
                "igst_amount",
                "total_amount",
                "final_offer_generated",
                "final_offer_pdf",
                "updated_at",
            ]
        )
        if old_pdf_name and old_pdf_name != quotation.final_offer_pdf.name:
            quotation.final_offer_pdf.storage.delete(old_pdf_name)

        return Response(
            QuotationSerializer(quotation, context={"request": request}).data,
            status=status.HTTP_200_OK,
        )



# create so from quotation

def _sales_order_pdf(sales_order):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    y = height - 48

    def write_line(text, font="Helvetica", size=10, gap=16):
        nonlocal y
        if y < 55:
            pdf.showPage()
            y = height - 48
        pdf.setFont(font, size)
        pdf.drawString(48, y, str(text)[:110])
        y -= gap

    write_line("SALES ORDER", "Helvetica-Bold", 18, 24)
    write_line(f"Sales Order: {sales_order.so_number}", "Helvetica-Bold")
    write_line(f"Date: {sales_order.so_date}")
    write_line(f"Quotation: {sales_order.quotation.quotation_number}")
    write_line(f"Company: {sales_order.company.name}")
    for value in (
        sales_order.company.address,
        ", ".join(filter(None, [
            sales_order.company.city.name if sales_order.company.city else "",
            sales_order.company.state.name if sales_order.company.state else "",
            sales_order.company.pincode,
        ])),
    ):
        if value:
            for line in wrap(value, 90):
                write_line(line)
    write_line("Items", "Helvetica-Bold", 12, 20)
    for index, item in enumerate(sales_order.quotation.items.all(), start=1):
        write_line(f"{index}. {item.module_name or item.description} — {item.amount}")
        if item.requirements:
            for line in wrap(item.requirements, 90):
                write_line(line)
    write_line(f"Total: {sales_order.total_amount}", "Helvetica-Bold", 12)
    if sales_order.delivery_terms:
        write_line("Delivery terms", "Helvetica-Bold", 11)
        for line in wrap(sales_order.delivery_terms, 90):
            write_line(line)
    if sales_order.payment_terms:
        write_line("Payment terms", "Helvetica-Bold", 11)
        for line in wrap(sales_order.payment_terms, 90):
            write_line(line)
    if sales_order.remarks:
        write_line("Remarks", "Helvetica-Bold", 11)
        for line in wrap(sales_order.remarks, 90):
            write_line(line)
    pdf.save()
    return buffer.getvalue()


def _refresh_sales_order_pdf(sales_order):
    previous_name = sales_order.so_pdf.name if sales_order.so_pdf else None
    sales_order.so_pdf.save(
        f"{sales_order.so_number}.pdf",
        ContentFile(_sales_order_pdf(sales_order)),
        save=False,
    )
    sales_order.save(update_fields=["so_pdf", "updated_at"])
    if previous_name and previous_name != sales_order.so_pdf.name:
        sales_order.so_pdf.storage.delete(previous_name)




@api_view(["POST"])
@permission_classes([IsAuthenticated])
@transaction.atomic
def create_sales_order_from_quotation(request, quotation_id):
    quotation = get_object_or_404(
        Quotation.objects.select_for_update().select_related(
            "company", "opportunity__company", "opportunity__lead__company"
        ),
        pk=quotation_id,
    )
    if not quotation.cpo_created:
        return Response(
            {"detail": "Mark the quotation for CPO processing before creating a sales order."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    if quotation.sales_orders.exists():
        return Response(
            {"detail": "Sales order already exists for this quotation"},
            status=status.HTTP_400_BAD_REQUEST,
        )

    company = quotation.company
    if company is None and quotation.opportunity:
        company = quotation.opportunity.company
        if company is None:
            try:
                company = quotation.opportunity.lead.company
            except (Lead.DoesNotExist, AttributeError):
                company = None
    if company is None:
        return Response(
            {"detail": "A company must be linked to the quotation before creating a sales order."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    current_year = timezone.localdate().year
    last_so = (
        SalesOrder.objects.filter(so_number__startswith=f"SO-{current_year}-")
        .order_by("-so_number")
        .first()
    )
    last_number = 0
    if last_so:
        try:
            last_number = int(last_so.so_number.rsplit("-", 1)[-1])
        except ValueError:
            last_number = 0
    so_number = f"SO-{current_year}-{last_number + 1:03d}"

    sales_order = SalesOrder.objects.create(
        so_number=so_number,
        so_date=timezone.localdate(),
        quotation=quotation,
        company=company,
        subtotal=quotation.subtotal,
        discount_amount=quotation.discount,
        taxable_amount=quotation.taxable_amount,
        tax_percentage=quotation.tax_percentage,
        tax_amount=quotation.tax_amount,
        total_amount=quotation.total_amount,
        tax_type=quotation.tax_type,
        cgst_amount=quotation.cgst_amount,
        sgst_amount=quotation.sgst_amount,
        igst_amount=quotation.igst_amount,
        packing_charges=quotation.packing_charges,
        payment_terms=request.data.get("payment_terms", ""),
        delivery_terms=request.data.get("delivery_terms", ""),
        remarks=request.data.get("remarks", ""),
        status="DRAFT",
        created_by=request.user,
    )
    _refresh_sales_order_pdf(sales_order)
    quotation.status = "CONVERTED"
    quotation.save(update_fields=["status", "updated_at"])

    return Response(
        {"so_number": sales_order.so_number, "id": sales_order.id},
        status=status.HTTP_201_CREATED,
    )



class SalesOrderListView(generics.ListAPIView):
    serializer_class = SalesOrderSerializer
    permission_classes = [IsAuthenticated]
    queryset = (
        SalesOrder.objects.select_related(
            "company", "company__city", "company__state", "company__country",
            "quotation", "created_by",
        )
        .prefetch_related("quotation__items__module", "company__contacts")
        .all()
    )


class SalesOrderDetailView(generics.RetrieveUpdateAPIView):
    serializer_class = SalesOrderSerializer
    permission_classes = [IsAuthenticated]
    queryset = (
        SalesOrder.objects.select_related(
            "company", "company__city", "company__state", "company__country",
            "quotation", "created_by",
        )
        .prefetch_related("quotation__items__module", "company__contacts")
    )

    def perform_update(self, serializer):
        sales_order = serializer.save()
        _refresh_sales_order_pdf(sales_order)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def get_order_dashboard(request):
    period = request.query_params.get("period", "yearly")
    try:
        year = int(request.query_params.get("year", timezone.localdate().year))
        if not 1 <= year <= 9998:
            raise ValueError
    except (TypeError, ValueError):
        return Response(
            {"detail": "year must be a valid year between 1 and 9998."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if period == "custom":
        start_date = parse_date(request.query_params.get("start_date", ""))
        end_date = parse_date(request.query_params.get("end_date", ""))
        if not start_date or not end_date:
            return Response(
                {"detail": "Custom periods require valid start_date and end_date values."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if start_date > end_date:
            return Response(
                {"detail": "start_date must be on or before end_date."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        period_label = f"{start_date} to {end_date}"
    elif period == "quarterly":
        try:
            quarter = int(request.query_params.get("quarter", 1))
            if quarter not in range(1, 5):
                raise ValueError
        except (TypeError, ValueError):
            return Response(
                {"detail": "quarter must be an integer from 1 to 4."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        start_month = (quarter - 1) * 3 + 1
        end_month = start_month + 2
        start_date = date(year, start_month, 1)
        end_date = date(year, end_month, monthrange(year, end_month)[1])
        period_label = f"Q{quarter} {year}"
    elif period == "half_yearly":
        try:
            half = int(request.query_params.get("half", 1))
            if half not in (1, 2):
                raise ValueError
        except (TypeError, ValueError):
            return Response(
                {"detail": "half must be either 1 or 2."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        start_month, end_month = (1, 6) if half == 1 else (7, 12)
        start_date = date(year, start_month, 1)
        end_date = date(year, end_month, monthrange(year, end_month)[1])
        period_label = f"H{half} {year}"
    elif period == "yearly":
        start_date = date(year, 1, 1)
        end_date = date(year, 12, 31)
        period_label = f"Annual {year}"
    else:
        return Response(
            {"detail": "period must be yearly, quarterly, half_yearly, or custom."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    quotation_queryset = Quotation.objects.filter(
        quotation_date__gte=start_date,
        quotation_date__lte=end_date,
    )
    final_offer_queryset = quotation_queryset.filter(final_offer_generated=True)
    cpo_queryset = quotation_queryset.filter(cpo_date__isnull=False)
    sales_order_queryset = SalesOrder.objects.filter(
        so_date__gte=start_date,
        so_date__lte=end_date,
    )

    def total_amount(queryset):
        return float(queryset.aggregate(total=Sum("total_amount"))["total"] or 0)

    def conversion_rate(converted, total):
        return round(converted * 100 / total, 2) if total else 0

    def monthly_aggregates(queryset, date_field):
        aggregates = {}
        for row in (
            queryset.annotate(month=TruncMonth(date_field))
            .values("month")
            .annotate(count=Count("id"), amount=Sum("total_amount"))
            .order_by("month")
        ):
            month = row["month"]
            month_key = month.date() if isinstance(month, datetime) else month
            aggregates[month_key] = {
                "count": row["count"],
                "amount": float(row["amount"] or 0),
            }
        return aggregates

    quotations_by_month = monthly_aggregates(quotation_queryset, "quotation_date")
    final_offers_by_month = monthly_aggregates(final_offer_queryset, "quotation_date")
    cpos_by_month = monthly_aggregates(cpo_queryset, "quotation_date")
    sales_orders_by_month = monthly_aggregates(sales_order_queryset, "so_date")

    monthly_history = []
    month_start = date(start_date.year, start_date.month, 1)
    while month_start <= end_date:
        month_end = date(
            month_start.year,
            month_start.month,
            monthrange(month_start.year, month_start.month)[1],
        )
        bucket_date = month_start
        quotation_month = quotations_by_month.get(bucket_date, {})
        final_offer_month = final_offers_by_month.get(bucket_date, {})
        cpo_month = cpos_by_month.get(bucket_date, {})
        sales_order_month = sales_orders_by_month.get(bucket_date, {})
        monthly_history.append(
            {
                "month": month_start.strftime("%b %Y"),
                "quotations": quotation_month.get("count", 0),
                "quotations_amount": quotation_month.get("amount", 0),
                "fo": final_offer_month.get("count", 0),
                "fo_amount": final_offer_month.get("amount", 0),
                "cpo": cpo_month.get("count", 0),
                "cpo_amount": cpo_month.get("amount", 0),
                "so": sales_order_month.get("count", 0),
                "so_amount": sales_order_month.get("amount", 0),
            }
        )
        if month_end >= end_date:
            break
        next_month = month_start.month + 1
        month_start = date(
            month_start.year + (1 if next_month > 12 else 0),
            1 if next_month > 12 else next_month,
            1,
        )

    quotation_count = quotation_queryset.count()
    final_offer_count = final_offer_queryset.count()
    cpo_count = cpo_queryset.count()
    sales_order_count = sales_order_queryset.count()
    sales_orders_from_cpo = sales_order_queryset.filter(
        quotation__cpo_date__isnull=False
    ).count()

    return Response(
        {
            "period": period_label,
            "start_date": str(start_date),
            "end_date": str(end_date),
            "summary": {
                "total_quotations": quotation_count,
                "quotations_amount": total_amount(quotation_queryset),
                "fo_generated": final_offer_count,
                "fo_amount": total_amount(final_offer_queryset),
                "cpo_received": cpo_count,
                "cpo_amount": total_amount(cpo_queryset),
                "total_so": sales_order_count,
                "so_amount": total_amount(sales_order_queryset),
                "conversion_rates": {
                    "quotation_to_fo": conversion_rate(final_offer_count, quotation_count),
                    "quotation_to_cpo": conversion_rate(cpo_count, quotation_count),
                    "cpo_to_sales_order": conversion_rate(sales_orders_from_cpo, cpo_count),
                },
            },
            "monthly_history": monthly_history,
        }
    )


class EmailLogListView(APIView):
    def get(self, request):
        logs = MessageLog.objects.filter(channel=Channel.EMAIL).select_related(
            "batch_lead", "template"
        ).order_by("-sent_at")
        limit = request.query_params.get("limit")
        if limit is not None:
            try:
                limit = min(max(int(limit), 1), 100)
            except ValueError:
                return Response(
                    {"detail": "limit must be a positive integer."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            logs = logs[:limit]
        serializer = MessageLogSerializer(logs, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)


class EmailTemplateViewSet(viewsets.ModelViewSet):
    queryset = MessageTemplate.objects.all()
    serializer_class = MessageTemplateSerializer

    def perform_create(self, serializer):
        serializer.save()

    #  logger.info(f"Email template '{serializer.instance.name}' created by user '{self.request.user.username}'.")


class EmailSequenceViewSet(viewsets.ModelViewSet):
    queryset = Sequence.objects.all()
    serializer_class = SequenceSerializer

    def perform_create(self, serializer):
        serializer.save()
        # logger.info(f"Email sequence '{serializer.instance.name}' created by user '{self.request.user.username}'.")


class EmailSequenceStepViewSet(viewsets.ModelViewSet):
    queryset = SequenceStep.objects.all()
    serializer_class = SequenceStepSerializer

    def perform_create(self, serializer):
        serializer.save()
        # logger.info(f"Email sequence step '{serializer.instance.step_number}' for sequence '{serializer.instance.sequence.name}' created by user '{self.request.user.username}'.")


class BatchViewSet(viewsets.ModelViewSet):
    queryset = Batch.objects.all()
    serializer_class = BatchSerializer

    def perform_create(self, serializer):
        batch = serializer.save()
        self._automatic_email_result = None
        transaction.on_commit(
            lambda batch_id=batch.pk: setattr(
                self,
                "_automatic_email_result",
                self._process_immediate_email_campaign(batch_id),
            )
        )

    def create(self, request, *args, **kwargs):
        response = super().create(request, *args, **kwargs)
        automatic_email_result = getattr(self, "_automatic_email_result", None)
        if automatic_email_result is not None:
            response.data["automatic_email"] = automatic_email_result
        return response

    @staticmethod
    def _process_immediate_email_campaign(batch_id):
        try:
            batch = (
                Batch.objects.select_related("sequence")
                .prefetch_related("sequence__steps")
                .get(pk=batch_id)
            )
        except Batch.DoesNotExist:
            logger.warning(
                "Skipping immediate campaign processing: batch %s no longer exists.",
                batch_id,
            )
            return None

        if batch.status != Batch.Status.ACTIVE:
            logger.info(
                "Skipping immediate processing for batch %s: status is %s.",
                batch.pk,
                batch.status,
            )
            return None
        if not batch.sequence_id:
            logger.info(
                "Skipping immediate processing for batch %s: no sequence is assigned.",
                batch.pk,
            )
            return None
        if batch.sequence.channel != Channel.EMAIL:
            logger.info(
                "Skipping immediate processing for batch %s: channel is %s.",
                batch.pk,
                batch.sequence.channel,
            )
            return None

        first_step = batch.sequence.steps.order_by("step_number").first()
        if first_step is None:
            logger.warning(
                "Skipping immediate processing for batch %s: sequence %s has no steps.",
                batch.pk,
                batch.sequence_id,
            )
            return None
        if first_step.delay_days != 0 or first_step.delay_minutes != 0:
            logger.info(
                "Skipping immediate processing for batch %s: first step delay is "
                "%s days and %s minutes.",
                batch.pk,
                first_step.delay_days,
                first_step.delay_minutes,
            )
            return None

        output = StringIO()
        try:
            call_command(
                "process_email_campaign",
                batch_id=batch.pk,
                stdout=output,
            )
        except Exception:
            logger.exception(
                "Immediate email campaign processing failed for batch %s.",
                batch.pk,
            )
            raise

        command_output = output.getvalue()
        sent_count = command_output.count("Sent to ")
        failed_count = command_output.count("Failed ")
        logger.info(
            "Immediate email campaign processing finished for batch %s: %s",
            batch.pk,
            command_output.strip() or "no command output",
        )
        return {"sent": sent_count, "failed": failed_count}

    @action(detail=True, methods=["get"], url_path="leads")
    def get_batch_leads(self, request, pk=None):
        batch = self.get_object()
        leads = batch.batch_leads.all()
        serializer = BatchLeadSerializer(leads, many=True)
        return Response(serializer.data)

    def _send_now(self, batch, template_id=None, enquiry_date=None, batch_lead=None):
        if batch.status != Batch.Status.ACTIVE:
            return Response(
                {"detail": "Resume this campaign before sending email."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if batch.sequence_id and batch.sequence.channel != Channel.EMAIL:
            return Response(
                {"detail": "Send now is only available for email campaigns."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        template = None
        if template_id is not None:
            template = get_object_or_404(MessageTemplate, pk=template_id)
            if template.channel != Channel.EMAIL:
                return Response(
                    {"detail": "Choose an email template."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        if not batch.sequence_id:
            if template is None:
                return Response(
                    {"detail": "Choose an email template before sending."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            sequence = Sequence.objects.create(
                name=f"{batch.name} - One-time email",
                channel=Channel.EMAIL,
            )
            SequenceStep.objects.create(
                sequence=sequence,
                step_number=1,
                template=template,
            )
            batch.sequence = sequence
            batch.save(update_fields=["sequence"])

        output = StringIO()
        command_options = {"batch_id": batch.pk, "force": True, "stdout": output}
        if template is not None:
            command_options["template_id"] = template.pk
        if enquiry_date:
            command_options["enquiry_date"] = enquiry_date.isoformat()
        if batch_lead is not None:
            command_options["batch_lead_id"] = batch_lead.pk
        call_command("process_email_campaign", **command_options)

        result = output.getvalue()
        sent_count = result.count("Sent to ")
        failed_count = result.count("Failed ")
        if not sent_count:
            detail = result.strip() or "No eligible email was sent."
            return Response(
                {"detail": detail, "sent": 0, "failed": failed_count},
                status=status.HTTP_400_BAD_REQUEST,
            )
        return Response(
            {"sent": sent_count, "failed": failed_count},
            status=status.HTTP_200_OK,
        )

    @action(detail=True, methods=["post"], url_path="send-now")
    def send_now(self, request, pk=None):
        batch = self.get_object()
        template_id = request.data.get("template")
        if not template_id and not batch.sequence_id:
            return Response(
                {"detail": "Choose an email template before sending."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if template_id:
            try:
                template_id = int(template_id)
            except (TypeError, ValueError):
                return Response(
                    {"template": "Choose a valid email template."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
        enquiry_date_value = request.data.get("enquiry_date")
        enquiry_date = None
        if enquiry_date_value:
            if not isinstance(enquiry_date_value, str):
                return Response(
                    {"enquiry_date": "Enter a valid date."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            enquiry_date = parse_date(enquiry_date_value)
            if enquiry_date is None:
                return Response(
                    {"enquiry_date": "Enter a valid date."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
        return self._send_now(
            batch,
            template_id=template_id,
            enquiry_date=enquiry_date,
        )

    @action(
        detail=True,
        methods=["post"],
        url_path=r"leads/(?P<lead_id>\d+)/send-now",
        url_name="send-lead-now",
    )
    def send_lead_now(self, request, pk=None, lead_id=None):
        batch = self.get_object()
        batch_lead = get_object_or_404(batch.batch_leads, lead_id=lead_id)
        return self._send_now(batch, batch_lead=batch_lead)

    @action(detail=True, methods=["patch"], url_path="pause", url_name="pause")
    def pause_batch(self, request, pk=None):
        batch = self.get_object()
        batch.status = Batch.Status.PAUSED
        batch.save()
        return Response({"status": "Batch paused."}, status=status.HTTP_200_OK)

    @action(detail=True, methods=["patch"], url_path="resume")
    def resume(self, request, pk=None):
        batch = self.get_object()
        batch.status = Batch.Status.ACTIVE
        batch.save(update_fields=["status"])
        automatic_email_result = []
        transaction.on_commit(
            lambda batch_id=batch.pk: automatic_email_result.append(
                self._process_immediate_email_campaign(batch_id)
            )
        )
        response_data = {"status": batch.status}
        if automatic_email_result and automatic_email_result[0] is not None:
            response_data["automatic_email"] = automatic_email_result[0]
        return Response(response_data, status=status.HTTP_200_OK)


# webhook


@csrf_exempt
def campaign_unsubscribe(request, token):
    if request.method not in ("GET", "POST"):
        return HttpResponse("Method not allowed.", status=405)

    batch_lead = resolve_unsubscribe_token(token)
    if batch_lead is None:
        return HttpResponse(
            "This unsubscribe link is invalid or no longer valid.", status=404
        )

    if request.method == "GET":
        return render(
            request,
            "campaigns/unsubscribe.html",
            {
                "already_unsubscribed": (
                    batch_lead.unsubscribed or is_email_suppressed(batch_lead.email)
                ),
                "completed": False,
                "mailgun_sync_failed": False,
            },
        )

    record_unsubscribe(batch_lead.email)
    mailgun_sync_failed = False
    try:
        sync_mailgun_suppression(batch_lead.email)
    except (ImproperlyConfigured, requests.RequestException):
        logger.exception("Could not sync an unsubscribe to Mailgun.")
        mailgun_sync_failed = True

    return render(
        request,
        "campaigns/unsubscribe.html",
        {
            "already_unsubscribed": False,
            "completed": True,
            "mailgun_sync_failed": mailgun_sync_failed,
        },
    )


def _is_valid_mailgun_signature(signature_data):
    signing_key = getattr(settings, "ANYMAIL", {}).get(
        "MAILGUN_WEBHOOK_SIGNING_KEY", ""
    )
    timestamp = str(signature_data.get("timestamp", ""))
    token = str(signature_data.get("token", ""))
    signature = str(signature_data.get("signature", ""))

    if not signing_key or not timestamp or not token or not signature:
        return False

    expected_signature = hmac.new(
        signing_key.encode(),
        f"{timestamp}{token}".encode(),
        hashlib.sha256,
    ).hexdigest()
    return hmac.compare_digest(signature, expected_signature)


@csrf_exempt
def mailgun_webhook(request):

    if request.method != "POST":
        return JsonResponse({"error": "POST method required"}, status=405)

    try:
        data = json.loads(request.body)
        if not _is_valid_mailgun_signature(data.get("signature", {})):
            return JsonResponse({"error": "Webhook signature is invalid."}, status=403)

        event_data = data.get("event-data", {})

        event = event_data.get("event")
        recipient = event_data.get("recipient", "").strip().lower()

        print("Mailgun Event:", event)
        print("Recipient:", recipient)

        # Mailgun event → MessageLog status
        event_status_map = {
            "accepted": MessageLog.Status.ACCEPTED,
            "delivered": MessageLog.Status.DELIVERED,
            "opened": MessageLog.Status.OPENED,
            "clicked": MessageLog.Status.CLICKED,
            "failed": MessageLog.Status.FAILED,
            "complained": MessageLog.Status.COMPLAINED,
            "unsubscribed": MessageLog.Status.UNSUBSCRIBED,
        }

        new_status = event_status_map.get(event)

        if not new_status:
            return JsonResponse(
                {"status": "ignored", "event": event, "message": "Event is not handled"}
            )

        if event == "unsubscribed":
            if not recipient:
                return JsonResponse(
                    {"error": "Unsubscribe event has no recipient."}, status=400
                )
            record_unsubscribe(recipient)

        message_id = (
            event_data.get("message", {}).get("headers", {}).get("message-id", "")
        )
        message_id = message_id.strip().strip("<>").lower()
        email_log = (
            MessageLog.objects.filter(
                channel=Channel.EMAIL,
                provider_message_id=message_id,
            ).first()
            if message_id
            else None
        )

        # Fall back to recipient for legacy events without a message ID.
        if not email_log:
            email_log = (
                MessageLog.objects.filter(
                    channel=Channel.EMAIL, to_address__iexact=recipient
                )
                .order_by("-sent_at")
                .first()
            )

        if not email_log:
            if event == "unsubscribed":
                return JsonResponse(
                    {"status": "suppressed", "event": event}
                )

            logger.warning("MessageLog not found for Mailgun recipient %s", recipient)
            return JsonResponse(
                {"status": "not_found", "event": event, "recipient": recipient}
            )

        # Update status
        email_log.status = new_status
        email_log.save(update_fields=["status"])

        print(f"MessageLog {email_log.id} updated: " f"{recipient} → {new_status}")

        return JsonResponse(
            {
                "status": "updated",
                "event": event,
                "recipient": recipient,
                "email_log_id": email_log.id,
                "message_log_id": email_log.id,
                "message_status": new_status,
            }
        )

    except json.JSONDecodeError:
        return JsonResponse(
            {"status": "error", "error": "Invalid JSON payload"}, status=400
        )

    except Exception as e:
        logger.exception("Mailgun webhook failed.")

        return JsonResponse(
            {
                "status": "error",
                "error": "Mailgun webhook processing failed.",
            },
            status=500,
        )


@csrf_exempt
def mailgun_inbound_webhook(request):
    if request.method != "POST":
        return JsonResponse({"error": "POST method required"}, status=405)

    signature_data = {
        "timestamp": request.POST.get("timestamp", ""),
        "token": request.POST.get("token", ""),
        "signature": request.POST.get("signature", ""),
    }
    if not _is_valid_mailgun_signature(signature_data):
        return JsonResponse({"error": "Webhook signature is invalid."}, status=403)

    try:
        sender = request.POST.get("sender", "").strip().lower()
        subject = request.POST.get("subject", "")
        body = request.POST.get("body-plain", "")

        print("Mailgun Reply Received")
        print("Sender:", sender)
        print("Subject:", subject)
        print("Body:", body)

        # Find the lead who replied
        batch_lead = BatchLead.objects.filter(email__iexact=sender).first()

        if not batch_lead:
            return JsonResponse(
                {
                    "status": "ignored",
                    "message": "No matching lead found",
                    "sender": sender,
                }
            )

        # Mark lead as replied
        batch_lead.replied = True
        batch_lead.replied_at = timezone.now()

        # Stop future emails
        batch_lead.save(update_fields=["replied", "replied_at"])

        # Update latest email log
        latest_log = (
            MessageLog.objects.filter(
                batch_lead=batch_lead,
                channel=Channel.EMAIL,
            )
            .order_by("-sent_at")
            .first()
        )

        if latest_log:
            latest_log.status = MessageLog.Status.REPLIED
            latest_log.save(update_fields=["status"])

        return JsonResponse(
            {
                "status": "updated",
                "sender": sender,
                "batch_lead_id": batch_lead.id,
                "email_log_id": latest_log.id if latest_log else None,
                "message_log_id": latest_log.id if latest_log else None,
            }
        )

    except Exception as e:
        return JsonResponse({"status": "error", "error": str(e)}, status=400)


@csrf_exempt
def whatsapp_webhook(request):
    verify_token = getattr(settings, "WHATSAPP_VERIFY_TOKEN", "")

    if request.method == "GET":
        mode = request.GET.get("hub.mode", "")
        token = request.GET.get("hub.verify_token", "")
        challenge = request.GET.get("hub.challenge", "")
        if (
            mode == "subscribe"
            and verify_token
            and hmac.compare_digest(token, verify_token)
        ):
            return HttpResponse(challenge, content_type="text/plain")
        return JsonResponse({"error": "Webhook verification failed."}, status=403)

    if request.method != "POST":
        return JsonResponse({"error": "GET or POST method required."}, status=405)

    app_secret = getattr(settings, "WHATSAPP_APP_SECRET", "")
    signature = request.headers.get("X-Hub-Signature-256", "")
    expected_signature = (
        "sha256="
        + hmac.new(app_secret.encode(), request.body, hashlib.sha256).hexdigest()
        if app_secret
        else ""
    )
    if not expected_signature or not hmac.compare_digest(signature, expected_signature):
        return JsonResponse({"error": "Webhook signature is invalid."}, status=403)

    try:
        payload = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({"error": "Invalid JSON payload."}, status=400)

    status_map = {
        "sent": MessageLog.Status.SENT,
        "delivered": MessageLog.Status.DELIVERED,
        "read": MessageLog.Status.READ,
        "failed": MessageLog.Status.FAILED,
    }
    status_rank = {
        MessageLog.Status.SENT: 1,
        MessageLog.Status.FAILED: 2,
        MessageLog.Status.DELIVERED: 3,
        MessageLog.Status.READ: 4,
    }
    updated = 0
    replied = 0

    for entry in payload.get("entry", []):
        for change in entry.get("changes", []):
            value = change.get("value", {})
            for provider_status in value.get("statuses", []):
                provider_message_id = provider_status.get("id", "")
                new_status = status_map.get(provider_status.get("status"))
                if not provider_message_id or not new_status:
                    continue
                message_log = MessageLog.objects.filter(
                    channel=Channel.WHATSAPP,
                    provider_message_id=provider_message_id,
                ).first()
                if not message_log:
                    continue
                if status_rank[new_status] < status_rank.get(message_log.status, 0):
                    continue
                errors = provider_status.get("errors", [])
                message_log.status = new_status
                message_log.error = (
                    (errors[0].get("title") or errors[0].get("message", ""))
                    if new_status == MessageLog.Status.FAILED and errors
                    else ""
                )
                message_log.save(update_fields=["status", "error"])
                updated += 1

            for inbound_message in value.get("messages", []):
                sender = re.sub(r"\D", "", inbound_message.get("from", ""))
                if not sender:
                    continue
                batch_lead = (
                    BatchLead.objects.filter(
                        phone__in=[sender, f"+{sender}"],
                        replied=False,
                        message_logs__channel=Channel.WHATSAPP,
                    )
                    .distinct()
                    .first()
                )
                if not batch_lead:
                    continue

                try:
                    replied_at = datetime.fromtimestamp(
                        float(inbound_message.get("timestamp")),
                        tz=dt_timezone.utc,
                    )
                except (TypeError, ValueError, OSError):
                    replied_at = timezone.now()
                batch_lead.replied = True
                batch_lead.replied_at = replied_at
                batch_lead.save(update_fields=["replied", "replied_at"])

                latest_log = (
                    MessageLog.objects.filter(
                        batch_lead=batch_lead,
                        channel=Channel.WHATSAPP,
                    )
                    .order_by("-sent_at")
                    .first()
                )
                if latest_log:
                    latest_log.status = MessageLog.Status.REPLIED
                    latest_log.save(update_fields=["status"])
                replied += 1

    return JsonResponse({"status": "processed", "updated": updated, "replied": replied})


# convert to opportunity


@api_view(["POST"])
def lead_convert_to_opportunity(request, lead_id):
    lead = get_object_or_404(Lead, pk=lead_id, is_active=True)

    if lead.is_opportunity:
        return Response(
            {"error": "Lead is already converted to an opportunity."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    opportunity = Opportunity.objects.create(
        lead=lead,
        company=lead.company,
        source=lead.source,
        status="PROSPECTING",
        is_converted_from_lead=True,
        converted_at=timezone.now(),
        converted_by=request.user if request.user.is_authenticated else None,
    )

    lead.is_opportunity = True
    lead.save(update_fields=["is_opportunity", "updated_at"])

    return Response(
        {
            "message": "Lead converted to opportunity successfully",
            "opportunity_id": opportunity.id,
            "is_opportunity": lead.is_opportunity,
            "is_converted_from_lead": opportunity.is_converted_from_lead,
            "company_name": lead.company.name,
        },
        status=status.HTTP_201_CREATED,
    )

    # convert to quotaitn


def generate_quotation_number():
    year = timezone.now().year

    prefix = f"QT-{year}-"

    last_quotation = (
        Quotation.objects.filter(quotation_number__startswith=prefix)
        .order_by("-id")
        .first()
    )

    if last_quotation:
        try:
            last_number = int(last_quotation.quotation_number.replace(prefix, ""))
        except ValueError:
            last_number = 0
    else:
        last_number = 0

    return f"{prefix}{last_number + 1:04d}"


class NextQuotationNumberView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response({"quotation_number": generate_quotation_number()})


class ConvertOpportunityToQuotationView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request, pk):
        opportunity = (
            Opportunity.objects.prefetch_related("modules")
            .select_related("company", "lead")
            .filter(id=pk, is_active=True)
            .first()
        )

        if opportunity is None:
            return Response(
                {"detail": "Opportunity not found."}, status=status.HTTP_404_NOT_FOUND
            )

        modules = list(opportunity.modules.all().order_by("module_number"))
        if not modules:
            return Response(
                {"detail": "Cannot convert opportunity without modules."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if opportunity.quotations.filter(
            status__in=["DRAFT", "SENT", "APPROVED", "REVISION"]
        ).exists():
            return Response(
                {"detail": "An active quotation already exists for this opportunity."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        raw_items = request.data.get("items")
        if not isinstance(raw_items, list) or not raw_items:
            return Response(
                {"detail": "At least one module item is required."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        module_map = {module.id: module for module in modules}
        submitted_by_module = {}

        for index, raw_item in enumerate(raw_items, start=1):
            if not isinstance(raw_item, dict):
                return Response(
                    {"detail": f"Item {index} is invalid."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            module_id = raw_item.get("module")
            amount_value = raw_item.get("amount")

            if module_id in (None, ""):
                return Response(
                    {"detail": f"Item {index} is missing a module reference."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            try:
                module_id = int(module_id)
            except (TypeError, ValueError):
                return Response(
                    {"detail": f"Item {index} has an invalid module id."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if module_id not in module_map:
                return Response(
                    {
                        "detail": f"Module {module_id} does not belong to this opportunity."
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if module_id in submitted_by_module:
                return Response(
                    {"detail": f"Module {module_id} was submitted more than once."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if amount_value in (None, ""):
                return Response(
                    {"detail": f"Amount is required for module {module_id}."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            try:
                amount = Decimal(str(amount_value))
            except InvalidOperation:
                return Response(
                    {"detail": f"Amount for module {module_id} must be numeric."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if amount < 0:
                return Response(
                    {"detail": f"Amount for module {module_id} cannot be negative."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            submitted_by_module[module_id] = amount

        missing_modules = [
            module.id for module in modules if module.id not in submitted_by_module
        ]
        if missing_modules:
            return Response(
                {"detail": "All opportunity modules must be included in the request."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        company = opportunity.company
        company_id = request.data.get("company_id")
        if company_id not in (None, ""):
            try:
                company_id = int(company_id)
            except (TypeError, ValueError):
                return Response(
                    {"detail": "Select a valid company."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            company = Company.objects.filter(pk=company_id, is_active=True).first()
            if company is None:
                return Response(
                    {"detail": "The selected company was not found."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            if opportunity.company_id != company.id:
                opportunity.company = company
                opportunity.save(update_fields=["company", "updated_at"])
        elif opportunity.company is None:
            return Response(
                {"detail": "Select a company before converting."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        quotation = Quotation.objects.create(
            quotation_number=generate_quotation_number(),
            quotation_date=timezone.now().date(),
            opportunity=opportunity,
            company=company,
            project_name=(opportunity.project_name or "").strip(),
            status="DRAFT",
            discount=Decimal("0.00"),
            tax_type="",
            tax_percentage=Decimal("0.00"),
            packing_charges=Decimal("0.00"),
            terms_conditions="",
            delivery_timeline="",
            created_by=request.user,
        )

        subtotal = Decimal("0.00")
        for module in modules:
            amount = submitted_by_module[module.id]
            QuotationItem.objects.create(
                quotation=quotation,
                module=module,
                description=module.description or module.module_name,
                amount=amount,
            )
            subtotal += amount

        quotation.subtotal = subtotal
        quotation.taxable_amount = subtotal
        quotation.tax_amount = Decimal("0.00")
        quotation.cgst_amount = Decimal("0.00")
        quotation.sgst_amount = Decimal("0.00")
        quotation.igst_amount = Decimal("0.00")
        quotation.total_amount = subtotal
        quotation.save()

        serializer = QuotationSerializer(quotation)
        return Response(
            {
                "message": "Opportunity converted to quotation successfully.",
                "quotation": serializer.data,
            },
            status=status.HTTP_201_CREATED,
        )
